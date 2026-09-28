from pypdf import PdfReader
from docx import Document


def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)

    elif name.endswith(".docx"):
        doc = Document(uploaded_file)
        parts = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)

        text = "\n".join(parts)

    elif name.endswith(".doc"):
        raise ValueError("Old .doc files aren't supported. Please save it as .docx or PDF and re-upload.")

    else:
        raise ValueError("Only PDF or DOCX files are supported.")

    text = text.strip()
    if not text:
        raise ValueError("Couldn't read any text from this file. Is it a scanned/image PDF?")
    return text